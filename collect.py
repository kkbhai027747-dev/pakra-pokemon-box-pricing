"""Read the supplier's Pokemon XLSX quotations without executing cell formulas."""

import argparse
import hashlib
import io
import json
import re
import sys
import time
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_CEILING, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
CENT = Decimal("0.01")


def money(value):
    return format(Decimal(value).quantize(CENT, rounding=ROUND_HALF_UP), ".2f")


def parse_offer(raw):
    if not isinstance(raw, str) or raw.lstrip().startswith("="):
        return None
    normalized = unicodedata.normalize("NFKC", raw)
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    if not lines or "价表" in lines[0]:
        return None
    if len(lines) == 1:
        if not re.search(r"原膜|礼盒|盒装|整盒|单盒", lines[0]):
            return None
        name, last = lines[0], ""
    else:
        name, last = " ".join(lines[:-1]), lines[-1]
    packaging = "待确认"
    if not re.search(r"散包|单包|拆包|拆盒|整箱|无膜|拆膜", name):
        if "礼盒" in name:
            packaging = "礼盒"
        elif re.search(r"原膜|盒装|整盒|单盒", name):
            packaging = "原膜盒装" if "原膜" in name else "盒装"
    result = {
        "product": name,
        "packaging": packaging,
        "cost": None,
        "status": "待确认",
        "note": "",
        "raw": raw,
    }
    # A second numeric line is ambiguous, even if the last line is a valid price.
    if any(re.fullmatch(r"[¥￥]?\s*\d+(?:\.\d+)?\s*元?", line) for line in lines[:-1]):
        result["note"] = "同一格含多个价格，需人工核对"
    elif last in {"/", "-", "—", "暂无", "缺货", "无货", "待定", "询价"}:
        result.update(status="未报价", note="原表未提供数字报价；不代表库存为零")
    elif re.fullmatch(r"[¥￥]?\s*\d+(?:\.\d{1,2})?\s*元?", last):
        amount = Decimal(re.sub(r"[¥￥元\s]", "", last))
        if amount > 0:
            result.update(cost=money(amount), status="已报价")
        else:
            result["note"] = "价格必须大于零"
    elif not last:
        result["note"] = "缺少价格行或表格排版已改变，需人工核对"
    else:
        result["note"] = "价格不是单一正数，需人工核对"
    if packaging == "待确认":
        notes = [result["note"], "未明确属于盒装"]
        result.update(status="待确认", note="；".join(note for note in notes if note))
    return result


def read_offers(path):
    book = load_workbook(path, data_only=False, keep_links=False)
    offers = []
    has_title = False
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" or not isinstance(cell.value, str):
                        continue
                    if "宝可梦" in cell.value and "价表" in cell.value:
                        has_title = True
                    offer = parse_offer(cell.value)
                    if offer:
                        offer.update(sheet=sheet.title, cell=cell.coordinate)
                        hidden_row = sheet.row_dimensions.get(cell.row)
                        hidden_col = any(
                            dimension.hidden
                            and dimension.min <= cell.column <= dimension.max
                            for dimension in sheet.column_dimensions.values()
                        )
                        hidden = (
                            sheet.sheet_state != "visible"
                            or (hidden_row and hidden_row.hidden)
                            or hidden_col
                        )
                        if hidden:
                            offer.update(status="待确认", note="隐藏工作表、行或列中的商品，需人工核对")
                        offers.append(offer)
    finally:
        book.close()
    if not has_title:
        raise ValueError("未找到宝可梦价表标题；请核对输入文件和表格格式")
    if not offers or not any(offer["packaging"] != "待确认" for offer in offers):
        raise ValueError("未提取到商品；表格格式可能已改变，保留原有结果")
    counts = Counter(offer["product"].casefold() for offer in offers)
    for offer in offers:
        if counts[offer["product"].casefold()] > 1:
            offer.update(status="待确认", note="同名商品出现多次，需人工核对")
    return offers


def validate_pricing(pricing):
    numbers = {}
    for key in ("fee_rate", "shipping", "profit_value", "round_to"):
        try:
            value = Decimal(str(pricing[key]))
        except (KeyError, InvalidOperation) as exc:
            raise ValueError(f"无效定价参数：{key}") from exc
        if not value.is_finite() or value < 0:
            raise ValueError(f"定价参数必须为有限非负数：{key}")
        numbers[key] = value
    if numbers["fee_rate"] >= 1 or numbers["round_to"] <= 0:
        raise ValueError("手续费率必须小于 1，取整单位必须大于 0")
    if numbers["round_to"] % CENT:
        raise ValueError("取整单位最小为 0.01 元")
    if pricing.get("fee_base") not in {"cost", "sale"}:
        raise ValueError("fee_base 只能为 cost 或 sale")
    if pricing.get("profit_mode") not in {"fixed", "cost_rate"}:
        raise ValueError("profit_mode 只能为 fixed 或 cost_rate")
    if not isinstance(pricing.get("confirmed"), bool):
        raise ValueError("confirmed 必须为 true 或 false")
    return numbers


def price_offer(cost, pricing):
    values = validate_pricing(pricing)
    cost = Decimal(str(cost))
    if not cost.is_finite() or cost <= 0:
        raise ValueError("进货价必须大于零")
    rate = values["fee_rate"]
    shipping = Decimal(money(values["shipping"]))
    target_profit = values["profit_value"]
    if pricing["profit_mode"] == "cost_rate":
        target_profit *= cost
    target_profit = Decimal(money(target_profit))
    if pricing["fee_base"] == "cost":
        raw_price = cost + Decimal(money(cost * rate)) + shipping + target_profit
    else:
        raw_price = (cost + shipping + target_profit) / (1 - rate)
    step = values["round_to"]
    price = (raw_price / step).to_integral_value(rounding=ROUND_CEILING) * step
    fee_base = cost if pricing["fee_base"] == "cost" else price
    fee = Decimal(money(fee_base * rate))
    return dict(
        fee=money(fee),
        shipping=money(shipping),
        profit=money(price - cost - fee - shipping),
        suggested_price=money(price),
        pricing_status="规则已确认" if pricing["confirmed"] else "试算，规则待确认",
    )


def quotation_date(path):
    match = re.match(r"(?:(20\d{2})[-.年])?(\d{1,2})[-.月](\d{1,2})(?:日)?(?=\D|$)", path.name)
    if not match:
        return None
    year, month, day = match.groups()
    if not year:
        year = next((part[:4] for part in reversed(path.parts[:-1]) if re.fullmatch(r"20\d{2}-\d{2}", part)), None)
    if not year:
        return None
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def choose_source(source, config_dir):
    location = Path(source["path"])
    if not location.is_absolute():
        location = config_dir / location
    if source["mode"] == "file":
        if not location.is_file() or location.suffix.lower() != ".xlsx":
            raise ValueError(f"找不到 XLSX 文件：{location}")
        return location.resolve()
    if source["mode"] != "directory" or not location.is_dir():
        raise ValueError(f"找不到价表目录或 source.mode 无效：{location}")
    iterator = location.rglob if source.get("recursive", False) else location.glob
    candidates = [
        path for path in iterator(source["pattern"])
        if path.is_file()
        and path.suffix.lower() == ".xlsx"
        and not path.name.startswith("~$")
    ]
    if not candidates:
        raise ValueError(f"目录中没有匹配的价表：{location}")
    # Dated supplier files take precedence over modification times of older copies.
    dated = [p for p in candidates if quotation_date(p)]
    return max(dated or candidates, key=lambda p: (quotation_date(p) or date.min, p.stat().st_mtime_ns, p.name)).resolve()


def index_offers(rows):
    """Keep duplicate products visible for review without selecting a price."""
    counts = Counter(row["product"].casefold() for row in rows)
    indexed = {}
    for row in rows:
        key = row["product"].casefold()
        if counts[key] == 1:
            indexed[key] = row
        else:
            indexed[key] = {
                **row,
                "cost": None,
                "suggested_price": None,
                "status": "待确认",
                "duplicate": True,
            }
    return indexed


def changes_since(previous, offers):
    before = index_offers(previous)
    after = index_offers(offers)
    compared_fields = ("cost", "status", "packaging", "suggested_price", "duplicate")
    changes = []
    for key in sorted(before.keys() | after.keys()):
        old, new = before.get(key), after.get(key)
        if old and new and all(old.get(field) == new.get(field) for field in compared_fields):
            continue
        if old is None:
            label = "新增"
        elif new is None:
            label = "本版未列出"
        else:
            label = "报价或定价变化"
        if new and new.get("duplicate"):
            label = "重复商品待确认"
        changes.append({
            "product": (new or old)["product"],
            "change": label,
            "old_cost": (old or {}).get("cost"),
            "new_cost": (new or {}).get("cost"),
            "old_price": (old or {}).get("suggested_price"),
            "new_price": (new or {}).get("suggested_price"),
        })
    return changes


def add_sheet(book, title, columns, rows):
    sheet = book.create_sheet(title)
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append([row.get(key) for key, _ in columns])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243D56")
        cell.font = Font(color="FFFFFF", bold=True)
    numeric = {"cost", "fee", "shipping", "profit", "suggested_price", "old_cost", "new_cost", "old_price", "new_price"}
    for cells in sheet.iter_rows(min_row=2):
        for cell, (key, _) in zip(cells, columns):
            if key in numeric and cell.value is not None:
                cell.value = float(cell.value)
                cell.number_format = '0.00'
            elif isinstance(cell.value, str):
                cell.data_type = "s"  # Keep supplier content as text, including leading '='.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, (key, _) in enumerate(columns, 1):
        letter = sheet.cell(1, index).column_letter
        sheet.column_dimensions[letter].width = 36 if key in {"product", "note", "raw", "value"} else 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def save_excel(path, report):
    book = Workbook()
    book.remove(book.active)
    columns = [("product", "商品原名"), ("packaging", "包装"), ("cost", "表内报价 / 元"),
               ("fee", "手续费 / 元"), ("shipping", "快递费 / 元"),
               ("suggested_price", "计算售价 / 元"), ("profit", "按配置估算利润 / 元"),
               ("pricing_status", "定价状态"), ("sheet", "源工作表"), ("cell", "源单元格"), ("raw", "原文")]
    add_sheet(book, "盒装报价", columns, [r for r in report["offers"] if r["status"] == "已报价"])
    add_sheet(book, "缺价及待确认", [("product", "商品原名"), ("packaging", "包装"), ("status", "状态"),
                              ("note", "原因"), ("sheet", "源工作表"), ("cell", "源单元格"), ("raw", "原文")],
              [r for r in report["offers"] if r["status"] != "已报价"])
    add_sheet(book, "本次变动", [("product", "商品原名"), ("change", "变动"), ("old_cost", "上次报价"),
                             ("new_cost", "本次报价"), ("old_price", "上次计算售价"), ("new_price", "本次计算售价")], report["changes"])
    info = [{"key": key, "value": str(value)} for key, value in report["summary"].items()]
    info.extend({"key": key, "value": str(value)} for key, value in report["pricing"].items())
    info.append({"key": "口径说明", "value": "金额暂按人民币理解；原膜依本次用途归为盒装候选，原表文字未明确计价单位。脚本不鉴定真伪。利润仅扣除配置中的进货价、手续费和快递费。"})
    add_sheet(book, "采集说明", [("key", "项目"), ("value", "内容")], info)
    book.save(path)
    book.close()


def source_status(previous, digest):
    if previous is None:
        return "首次采集"
    if previous["source_sha256"] == digest:
        return "源文件未变化，仍沿用原报价"
    return "源文件已变化"


def collect(config, config_dir):
    validate_pricing(config["pricing"])
    source = choose_source(config["source"], config_dir)
    data = source.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    offers = read_offers(io.BytesIO(data))
    for offer in offers:
        if offer["status"] == "已报价":
            offer.update(price_offer(offer["cost"], config["pricing"]))
    out = Path(config.get("output_dir", "output"))
    if not out.is_absolute():
        out = config_dir / out
    out.mkdir(parents=True, exist_ok=True)
    latest_json = out / "latest.json"
    previous = json.loads(latest_json.read_text(encoding="utf-8")) if latest_json.exists() else None
    now = datetime.now().astimezone()
    source_date = quotation_date(source)
    summary = {
        "采集时间": now.isoformat(timespec="seconds"), "源文件": str(source), "源文件SHA256": digest,
        "源文件修改时间": datetime.fromtimestamp(source.stat().st_mtime).astimezone().isoformat(timespec="seconds"),
        "文件名所示报价日期": source_date.isoformat() if source_date else "未知",
        "报价距今天数": (now.date() - source_date).days if source_date else "未知",
        "源文件状态": source_status(previous, digest),
        "商品数": len(offers), "有报价盒装数": sum(r["status"] == "已报价" for r in offers),
        "缺价或待确认数": sum(r["status"] != "已报价" for r in offers),
        "定价状态": "规则已确认" if config["pricing"]["confirmed"] else "试算，规则待确认",
        "每日来源状态": "按当前配置读取本地价表；新文件须先保存到该目录",
    }
    report = dict(summary=summary, source_sha256=digest, pricing=config["pricing"], offers=offers,
                  changes=changes_since(previous["offers"], offers) if previous else [])
    history = out / "history" / now.strftime("%Y-%m-%d")
    history.mkdir(parents=True, exist_ok=True)
    stem = now.strftime("%H%M%S-%f")
    excel_path = history / f"{stem}.xlsx"
    json_path = history / f"{stem}.json"
    save_excel(excel_path, report)
    content = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    json_path.write_text(content, encoding="utf-8")
    # Only replace current results after the complete source has parsed successfully.
    excel_tmp = out / "latest.tmp.xlsx"
    excel_tmp.write_bytes(excel_path.read_bytes())
    excel_tmp.replace(out / "latest.xlsx")
    json_tmp = out / "latest.tmp.json"
    json_tmp.write_text(content, encoding="utf-8")
    json_tmp.replace(latest_json)
    print(json.dumps({**summary, "本次变化数": len(report["changes"]), "结果": str(out / "latest.xlsx")}, ensure_ascii=False, indent=2))
    return report


def main():
    parser = argparse.ArgumentParser(description="宝可梦盒装本地价表采集；不访问微信或供应商网站")
    parser.add_argument("--config", type=Path, default=ROOT / "config.json")
    parser.add_argument("--file", type=Path, help="本次读取指定 XLSX，覆盖配置中的目录来源")
    parser.add_argument("--daily", action="store_true", help="立即采集一次，随后每天指定时间采集；窗口须保持运行")
    args = parser.parse_args()
    while True:
        try:
            config = json.loads(args.config.read_text(encoding="utf-8-sig"))
            hour, minute = map(int, config.get("daily_at", "09:00").split(":"))
            if not (0 <= hour <= 23 and 0 <= minute <= 59):
                raise ValueError("daily_at 必须为 HH:MM")
            if args.file:
                config["source"] = {"mode": "file", "path": str(args.file.resolve())}
            collect(config, args.config.resolve().parent)
        except Exception as exc:
            print(f"采集失败：{exc}。未完成的采集不会作为有效报价发布。", file=sys.stderr, flush=True)
            if not args.daily:
                return 1
            # Retry after a minute, including when a file was still being copied.
            time.sleep(60)
            continue
        if not args.daily:
            return 0
        now = datetime.now()
        next_run = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
        if next_run <= now:
            next_run += timedelta(days=1)
        print(f"下次采集：{next_run:%Y-%m-%d %H:%M}（电脑本地时间）；请保持窗口运行。", flush=True)
        while datetime.now() < next_run:
            time.sleep(min(30, max(0, (next_run - datetime.now()).total_seconds())))


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("已停止每日采集。")
