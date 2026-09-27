"""Focused parser, workbook-safety, and pricing regression tests."""

from pathlib import Path
from contextlib import redirect_stdout
import io
import os
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from collect import changes_since, choose_source, collect, parse_offer, price_offer, read_offers


class OfferTests(unittest.TestCase):
    def test_original_sealed_box_quote(self):
        raw = "M5原膜\n538"
        offer = parse_offer(raw)
        self.assertIsNotNone(offer)
        self.assertIn("M5", offer["product"])
        self.assertEqual(offer["packaging"], "原膜盒装")
        self.assertEqual(offer["cost"], "538.00")
        self.assertEqual(offer["status"], "已报价")
        self.assertEqual(offer["raw"], raw)

    def test_multiline_gift_and_missing_quote(self):
        gift = parse_offer("宝可梦礼盒\n特殊版本\n1000")
        self.assertIsNotNone(gift)
        self.assertEqual(gift["packaging"], "礼盒")
        self.assertEqual(gift["cost"], "1000.00")
        self.assertIn("特殊版本", gift["product"])

        missing = parse_offer("火箭队礼盒\n/")
        self.assertIsNotNone(missing)
        self.assertIsNone(missing["cost"])
        self.assertEqual(missing["status"], "未报价")

    def test_title_is_not_a_product(self):
        self.assertIsNone(parse_offer("日版宝可梦销售价表"))

    def test_box_name_without_same_cell_price_is_retained_for_review(self):
        offer = parse_offer("M4原膜")
        self.assertIsNotNone(offer)
        self.assertEqual(offer["product"], "M4原膜")
        self.assertIsNone(offer["cost"])
        self.assertEqual(offer["status"], "待确认")

    def test_ambiguous_or_invalid_prices_are_not_quotes(self):
        for price_text in ("500\n600", "500-600", "500×2", "-500", "1,000"):
            with self.subTest(price_text=price_text):
                offer = parse_offer("M5原膜\n" + price_text)
                if offer is not None:
                    self.assertIsNone(offer["cost"])
                    self.assertNotEqual(offer["status"], "已报价")


class PricingTests(unittest.TestCase):
    def pricing(self, **changes):
        settings = {
            "confirmed": True,
            "fee_rate": "0.06",
            "fee_base": "cost",
            "shipping": "70",
            "profit_mode": "fixed",
            "profit_value": "200",
            "round_to": "1",
        }
        settings.update(changes)
        return settings

    def test_user_example_cost_based_fee(self):
        result = price_offer("1000", self.pricing())
        self.assertEqual(result["fee"], "60.00")
        self.assertEqual(result["shipping"], "70.00")
        self.assertEqual(result["profit"], "200.00")
        self.assertEqual(result["suggested_price"], "1330.00")

    def test_sale_based_fee_preserves_target_profit_after_rounding(self):
        result = price_offer("1000", self.pricing(fee_base="sale"))
        self.assertEqual(result["suggested_price"], "1352.00")
        self.assertEqual(result["fee"], "81.12")
        self.assertEqual(result["profit"], "200.88")

    def test_cost_percentage_profit(self):
        result = price_offer(
            "500", self.pricing(profit_mode="cost_rate", profit_value="0.20")
        )
        self.assertEqual(result["fee"], "30.00")
        self.assertEqual(result["profit"], "100.00")
        self.assertEqual(result["suggested_price"], "700.00")

    def test_unconfirmed_rule_is_marked_as_trial(self):
        result = price_offer("1000", self.pricing(confirmed=False))
        self.assertEqual(result["pricing_status"], "试算，规则待确认")

    def test_invalid_pricing_parameters_fail_explicitly(self):
        cases = [
            ("-1", {}),
            ("1000", {"fee_rate": "1"}),
            ("1000", {"fee_rate": "1.2"}),
            ("1000", {"round_to": "0"}),
            ("1000", {"round_to": "-1"}),
        ]
        for cost, changes in cases:
            with self.subTest(cost=cost, changes=changes):
                with self.assertRaises(ValueError):
                    price_offer(cost, self.pricing(**changes))


class WorkbookTests(unittest.TestCase):
    def read_book(self, workbook):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "supplier.xlsx"
            workbook.save(path)
            before = path.read_bytes()
            offers = read_offers(path)
            self.assertEqual(path.read_bytes(), before, "Source workbook was modified")
            return offers

    def test_all_columns_are_scanned_without_treating_formulas_as_quotes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "供应报价"
        sheet["A1"] = "日版宝可梦销售价表"
        sheet["L3"] = "M5原膜\n538"
        sheet["M3"] = '=DISPIMG("ID_REFERENCE",1)'
        sheet["N3"] = '=HYPERLINK("https://example.invalid/","M4原膜 568")'
        sheet["O3"] = "快递费\n70"
        sheet["P3"] = "不明对象\n900"
        sheet["A8"] = "M4原膜\n568"
        sheet.row_dimensions[8].hidden = True
        hidden = workbook.create_sheet("过期报价")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "M3原膜\n438"

        offers = self.read_book(workbook)
        by_location = {(item["sheet"], item["cell"]): item for item in offers}
        visible = by_location[("供应报价", "L3")]
        self.assertEqual(visible["cost"], "538.00")
        self.assertEqual(visible["status"], "已报价")
        for cell in ("A1", "M3", "N3"):
            self.assertNotIn(("供应报价", cell), by_location)
        for cell in ("O3", "P3"):
            item = by_location.get(("供应报价", cell))
            if item is not None:
                self.assertNotEqual(item["status"], "已报价")
        for location in (("供应报价", "A8"), ("过期报价", "A1")):
            item = by_location.get(location)
            if item is not None:
                self.assertEqual(item["status"], "待确认")

    def test_inline_rich_text_remains_readable(self):
        workbook = Workbook()
        workbook.active["A1"] = "日版宝可梦销售价表"
        workbook.active["C3"] = CellRichText(
            "M5", TextBlock(InlineFont(b=True), "原膜"), "\n538"
        )
        offers = self.read_book(workbook)
        self.assertEqual(len(offers), 1)
        self.assertEqual(offers[0]["cost"], "538.00")
        self.assertEqual(offers[0]["cell"], "C3")


class DailyCollectionTests(unittest.TestCase):
    def test_filename_date_wins_over_mtime_and_lock_files_are_ignored(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_dir = root / "2026-09"
            source_dir.mkdir()
            older = source_dir / "9.8- 日版宝可梦销售价表.xlsx"
            newer = source_dir / "9.21- 日版宝可梦销售价表.xlsx"
            lock_file = source_dir / "~$9.30- 日版宝可梦销售价表.xlsx"
            for path in (older, newer, lock_file):
                path.touch()
            os.utime(newer, (100, 100))
            os.utime(older, (200, 200))
            os.utime(lock_file, (300, 300))
            source = {"mode": "directory", "path": "2026-09", "pattern": "*.xlsx"}
            self.assertEqual(choose_source(source, root), newer.resolve())

            older.unlink()
            newer.unlink()
            with self.assertRaises(ValueError):
                choose_source(source, root)

    def test_missing_quote_and_product_omission_are_distinct_changes(self):
        previous = [parse_offer("M5原膜\n538"), parse_offer("M4原膜\n568")]
        current = [parse_offer("M5原膜\n/")]
        changes = {item["product"]: item for item in changes_since(previous, current)}
        missing_quote = changes["M5原膜"]
        self.assertEqual(missing_quote["change"], "报价或定价变化")
        self.assertEqual(missing_quote["old_cost"], "538.00")
        self.assertIsNone(missing_quote["new_cost"])
        omitted = changes["M4原膜"]
        self.assertEqual(omitted["change"], "本版未列出")
        self.assertEqual(omitted["old_cost"], "568.00")
        self.assertIsNone(omitted["new_cost"])

    def test_duplicate_product_is_flagged_and_recovery_is_a_change(self):
        single = [parse_offer("M5原膜\n538")]
        duplicates = [parse_offer("M5原膜\n538"), parse_offer("M5原膜\n568")]
        changes = changes_since(single, duplicates)
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]["product"], "M5原膜")
        self.assertEqual(changes[0]["change"], "重复商品待确认")
        self.assertIsNone(changes[0]["new_cost"])

        recovered = changes_since(duplicates, single)
        self.assertEqual(len(recovered), 1)
        self.assertEqual(recovered[0]["product"], "M5原膜")
        self.assertNotIn(recovered[0]["change"], ("新增", "本版未列出"))
        self.assertEqual(recovered[0]["new_cost"], "538.00")

    def test_unchanged_source_has_no_changes_and_failure_keeps_latest_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "supplier.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "日版宝可梦销售价表"
            workbook.active["C3"] = "M5原膜\n1000"
            workbook.save(source)
            config = {
                "source": {"mode": "file", "path": source.name},
                "output_dir": "output",
                "pricing": {
                    "confirmed": False,
                    "fee_rate": "0.06",
                    "fee_base": "cost",
                    "shipping": "70",
                    "profit_mode": "fixed",
                    "profit_value": "200",
                    "round_to": "1",
                },
            }
            with redirect_stdout(io.StringIO()):
                first = collect(config, root)
                second = collect(config, root)
            self.assertEqual(first["offers"][0]["suggested_price"], "1330.00")
            self.assertEqual(second["source_sha256"], first["source_sha256"])
            self.assertEqual(second["changes"], [])
            self.assertEqual(second["summary"]["源文件状态"], "源文件未变化，仍沿用原报价")
            latest_paths = [root / "output" / name for name in ("latest.json", "latest.xlsx")]
            previous_bytes = {path: path.read_bytes() for path in latest_paths}

            workbook.active["A1"] = "无关文档"
            workbook.save(source)
            with redirect_stdout(io.StringIO()), self.assertRaises(ValueError):
                collect(config, root)
            for path in latest_paths:
                self.assertEqual(path.read_bytes(), previous_bytes[path])


if __name__ == "__main__":
    unittest.main()
